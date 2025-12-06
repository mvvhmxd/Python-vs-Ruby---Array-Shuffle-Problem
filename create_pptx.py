from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# Create presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
DARK_BG = RGBColor(10, 10, 15)
PYTHON_BLUE = RGBColor(55, 118, 171)
RUBY_RED = RGBColor(204, 52, 45)
ACCENT = RGBColor(99, 102, 241)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(161, 161, 170)

def add_title_slide(title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK_BG
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.3), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.3), Inches(0.8))
        tf2 = sub_box.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(24)
        p2.font.color.rgb = GRAY
        p2.alignment = PP_ALIGN.CENTER
    return slide

def add_content_slide(title, bullets, color=ACCENT):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK_BG
    
    # Header bar
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.2))
    header.fill.solid()
    header.fill.fore_color.rgb = color
    header.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Bullets
    content_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(12), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + bullet
        p.font.size = Pt(24)
        p.font.color.rgb = WHITE
        p.space_before = Pt(12)
    return slide

def add_code_slide(title, python_code, ruby_code):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK_BG
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Python header
    py_header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.1), Inches(6), Inches(0.5))
    py_header.fill.solid()
    py_header.fill.fore_color.rgb = PYTHON_BLUE
    py_header.line.fill.background()
    py_label = slide.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(5), Inches(0.5))
    py_label.text_frame.paragraphs[0].text = "Python"
    py_label.text_frame.paragraphs[0].font.size = Pt(18)
    py_label.text_frame.paragraphs[0].font.bold = True
    py_label.text_frame.paragraphs[0].font.color.rgb = WHITE
    
    # Python code
    py_code_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.7), Inches(6), Inches(4.8))
    tf = py_code_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = python_code
    p.font.size = Pt(14)
    p.font.name = "Consolas"
    p.font.color.rgb = RGBColor(255, 212, 59)
    
    # Ruby header
    rb_header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.1), Inches(6), Inches(0.5))
    rb_header.fill.solid()
    rb_header.fill.fore_color.rgb = RUBY_RED
    rb_header.line.fill.background()
    rb_label = slide.shapes.add_textbox(Inches(7), Inches(1.15), Inches(5), Inches(0.5))
    rb_label.text_frame.paragraphs[0].text = "Ruby"
    rb_label.text_frame.paragraphs[0].font.size = Pt(18)
    rb_label.text_frame.paragraphs[0].font.bold = True
    rb_label.text_frame.paragraphs[0].font.color.rgb = WHITE
    
    # Ruby code
    rb_code_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.7), Inches(6), Inches(4.8))
    tf = rb_code_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = ruby_code
    p.font.size = Pt(14)
    p.font.name = "Consolas"
    p.font.color.rgb = RGBColor(255, 107, 157)
    return slide

def add_comparison_slide(title, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK_BG
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    # Table headers
    headers = ["Feature", "Python", "Ruby"]
    col_widths = [Inches(4), Inches(4.5), Inches(4.5)]
    x_positions = [Inches(0.5), Inches(4.5), Inches(9)]
    
    for i, (header, x) in enumerate(zip(headers, x_positions)):
        color = [ACCENT, PYTHON_BLUE, RUBY_RED][i]
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.2), col_widths[i], Inches(0.6))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = color
        header_box.line.fill.background()
        
        text_box = slide.shapes.add_textbox(x, Inches(1.25), col_widths[i], Inches(0.5))
        text_box.text_frame.paragraphs[0].text = header
        text_box.text_frame.paragraphs[0].font.size = Pt(20)
        text_box.text_frame.paragraphs[0].font.bold = True
        text_box.text_frame.paragraphs[0].font.color.rgb = WHITE
        text_box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Data rows
    y = 1.9
    for row in data:
        for i, (cell, x) in enumerate(zip(row, x_positions)):
            text_box = slide.shapes.add_textbox(x, Inches(y), col_widths[i], Inches(0.5))
            text_box.text_frame.paragraphs[0].text = cell
            text_box.text_frame.paragraphs[0].font.size = Pt(16)
            text_box.text_frame.paragraphs[0].font.color.rgb = WHITE
            text_box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        y += 0.55
    return slide

# ===================== CREATE SLIDES =====================

# Slide 1: Title
add_title_slide("Python vs Ruby", "Array Shuffle Problem\nCSE 451: Concepts of Programming Languages")

# Slide 2: Problem Statement
add_content_slide("The Problem", [
    "Given a string s and an integer array indices of the same length",
    "Shuffle string so character at position i moves to indices[i]",
    "",
    "Example:",
    "   Input: s = \"art\", indices = [1, 0, 2]",
    "   Output: \"rat\"",
    "",
    "   'a' at index 0 → goes to index 1",
    "   'r' at index 1 → goes to index 0", 
    "   't' at index 2 → stays at index 2"
])

# Slide 3: Python Introduction
add_content_slide("Python - Introduction", [
    "Created by Guido van Rossum in 1991",
    "Multi-paradigm: Imperative, Object-Oriented, Functional",
    "Interpreted language (CPython implementation)",
    "Dynamic typing - variables allocated at runtime",
    "Known for clean, readable syntax",
    "Huge ecosystem with pip package manager",
    "Used in web dev, data science, AI, automation"
], PYTHON_BLUE)

# Slide 4: Python Technical Details
add_content_slide("Python - Technical Details", [
    "Data Types:",
    "   • Simple: int, float, bool, str, NoneType",
    "   • Structured: list, tuple, dict, set, class objects",
    "",
    "Statement Types:",
    "   • for x in iterable, for i in range(n)",
    "   • while condition, if/elif/else",
    "   • try/except/finally, with statement",
    "",
    "Subprograms:",
    "   • Functions (def), Lambda expressions",
    "   • Methods in classes, Exception handlers",
    "   • Threading module for concurrency"
], PYTHON_BLUE)

# Slide 5: Python Strengths & Weaknesses
add_content_slide("Python - Strengths & Weaknesses", [
    "Strengths:",
    "   ✓ Easy to learn and read",
    "   ✓ Huge library ecosystem (NumPy, Django, etc.)",
    "   ✓ Great documentation and community",
    "   ✓ Cross-platform compatibility",
    "",
    "Weaknesses:",
    "   ✗ Slower than compiled languages",
    "   ✗ GIL limits true multi-threading",
    "   ✗ Mobile development not ideal",
    "   ✗ Runtime errors instead of compile-time"
], PYTHON_BLUE)

# Slide 6: Ruby Introduction
add_content_slide("Ruby - Introduction", [
    "Created by Yukihiro Matsumoto in 1995",
    "Object-Oriented (everything is an object!)",
    "Interpreted language (MRI/YARV implementation)",
    "Dynamic typing - flexible variable allocation",
    "Philosophy: \"Developer happiness\"",
    "Famous for Ruby on Rails framework",
    "Flexible syntax - multiple ways to do things"
], RUBY_RED)

# Slide 7: Ruby Technical Details
add_content_slide("Ruby - Technical Details", [
    "Data Types:",
    "   • Simple: Integer, Float, TrueClass, FalseClass, String, Symbol",
    "   • Structured: Array, Hash, Range, Class objects",
    "",
    "Statement Types:",
    "   • each, times, upto/downto loops",
    "   • while/until, if/elsif/else/unless",
    "   • begin/rescue/ensure for exceptions",
    "",
    "Subprograms:",
    "   • Methods (def), Blocks { } and do/end",
    "   • Procs and Lambdas",
    "   • Threads for concurrency"
], RUBY_RED)

# Slide 8: Ruby Strengths & Weaknesses
add_content_slide("Ruby - Strengths & Weaknesses", [
    "Strengths:",
    "   ✓ Elegant, expressive syntax",
    "   ✓ Everything is an object - pure OOP",
    "   ✓ Ruby on Rails for web development",
    "   ✓ Great for DSLs (Domain Specific Languages)",
    "",
    "Weaknesses:",
    "   ✗ Slower than Python in many cases",
    "   ✗ Smaller community than Python",
    "   ✗ Less popular outside web dev",
    "   ✗ Debugging can be challenging"
], RUBY_RED)

# Slide 9: Code Comparison
py_code = """def shuffle_string(s, indices):
    result = [''] * len(s)
    for i, char in enumerate(s):
        result[indices[i]] = char
    return ''.join(result)

# Usage:
shuffle_string("art", [1, 0, 2])
# Returns: "rat"
"""

rb_code = """def shuffle_string(s, indices)
  result = Array.new(s.length, '')
  s.each_char.with_index do |char, i|
    result[indices[i]] = char
  end
  result.join
end

# Usage:
shuffle_string("art", [1, 0, 2])
# Returns: "rat"
"""
add_code_slide("Code Comparison", py_code, rb_code)

# Slide 10: Syntax Comparison Table
comparison_data = [
    ["Language Type", "Multi-paradigm", "Object-Oriented"],
    ["Implementation", "Interpreted (CPython)", "Interpreted (MRI)"],
    ["Variable Allocation", "Dynamic", "Dynamic"],
    ["Array Creation", "[''] * len(s)", "Array.new(s.length, '')"],
    ["Iteration", "enumerate(s)", "s.each_char.with_index"],
    ["Block Syntax", "indentation", "do/end or { }"],
    ["Join Array", "''.join(list)", "array.join"],
]
add_comparison_slide("Language Comparison", comparison_data)

# Slide 11: Algorithm Visualization
add_content_slide("Algorithm Step-by-Step", [
    "Input: s = \"art\", indices = [1, 0, 2]",
    "Result array starts empty: [_, _, _]",
    "",
    "Step 1: s[0] = 'a' → result[indices[0]] = result[1] = 'a'",
    "        Result: [_, a, _]",
    "",
    "Step 2: s[1] = 'r' → result[indices[1]] = result[0] = 'r'",
    "        Result: [r, a, _]",
    "",
    "Step 3: s[2] = 't' → result[indices[2]] = result[2] = 't'",
    "        Result: [r, a, t]",
    "",
    "Final: \"rat\""
])

# Slide 12: Challenges
add_content_slide("Challenges Faced", [
    "Setting up Ruby on Windows:",
    "   • Had to install RubyInstaller with DevKit",
    "   • PATH configuration required",
    "",
    "Learning Ruby syntax:",
    "   • Block syntax (do/end) was new coming from Python",
    "   • each_char.with_index vs Python's enumerate()",
    "",
    "Testing:",
    "   • Verified both solutions produce identical output",
    "   • Created step-by-step visualizer for debugging"
])

# Slide 13: Demo slide with QR
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = DARK_BG

title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
tf = title_box.text_frame
p = tf.paragraphs[0]
p.text = "Try the Live Demo!"
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = WHITE
p.alignment = PP_ALIGN.CENTER

# Link
link_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12), Inches(0.6))
tf = link_box.text_frame
p = tf.paragraphs[0]
p.text = "https://mvvhmxd.github.io/Python-vs-Ruby---Array-Shuffle-Problem/"
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(99, 102, 241)
p.alignment = PP_ALIGN.CENTER

# QR Code
qr_path = os.path.join(os.path.dirname(__file__), "qrcode.jpg")
if os.path.exists(qr_path):
    slide.shapes.add_picture(qr_path, Inches(4.5), Inches(2.2), width=Inches(4.3))

# Instructions
inst_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12), Inches(0.6))
tf = inst_box.text_frame
p = tf.paragraphs[0]
p.text = "Scan QR code or visit the link to see the interactive algorithm visualizer"
p.font.size = Pt(18)
p.font.color.rgb = GRAY
p.alignment = PP_ALIGN.CENTER

# Slide 14: Conclusion
add_content_slide("Conclusion", [
    "Both languages successfully solve the problem",
    "Same algorithm, different syntax",
    "",
    "Time Complexity: O(n)",
    "Space Complexity: O(n)",
    "",
    "Python: More explicit and widely used",
    "Ruby: More flexible and expressive",
    "",
    "Key insight: Understanding concepts matters more than syntax!"
])

# Slide 15: Thank You
add_title_slide("Thank You!", "Questions?")

# Save
prs.save("CSE451_Presentation.pptx")
print("Presentation saved as CSE451_Presentation.pptx")
